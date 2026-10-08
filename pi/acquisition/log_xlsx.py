"""엑셀 기록 — 파이프라인(pi/pipeline.py)의 출력 하나.

Frame + Result(analysis/analyzer.py) 하나를 CSV 한 행으로 쓰고, 끝날 때(Ctrl+C / LOG_SECONDS / kill / SSH 끊김) 엑셀로 만든다.
열 구성은 HX_LABELS / IMU_LABELS 로 고정된다. 응답 없는 모듈, 못 연 IMU, 보정값이 없는 채널도
열은 그대로 두고 칸만 비운다 (요약 시트에 남긴다).

만들어지는 파일 (LOG_DIR/)
    <이름>.csv    기록하면서 1초마다 디스크에 쓴다. 도중에 전원이 나가도 여기까지 남고
                  엑셀로 바로 열린다 (UTF-8 BOM 이라 한글이 안 깨진다)
    <이름>.xlsx   끝날 때 CSV 로 만든다. 측정 + 요약 시트

CSV 만 남았을 때 엑셀 다시 만들기:
    python3 -c "from acquisition import log_xlsx; log_xlsx.build_xlsx('X.csv', 'X.xlsx')"
"""
import csv
import datetime
import os
import time

from acquisition.console import say, text_width
from analysis.analyzer import check_list
from analysis.risk import CAUTION, DANGER, OK, WARNING
from analysis.risk import name as risk_name

FLUSH_EVERY = 1.0           # CSV 를 디스크까지 내리는 주기(초)
MAX_EVENTS = 5000           # '판정' 시트에 적는 알림 변화 수 한계
EXCEL_MAX_ROWS = 1048576    # 엑셀 시트 하나의 행 한계
TIME_FMT = "%Y-%m-%d %H:%M:%S.%f"
PIP_HINT = "pip install openpyxl   (파이: sudo apt install python3-openpyxl)"

# (ImuReading 속성, CSV 열 이름 뒷부분, 소수 자리)
IMU_COLS = (("ax", "ax_G", 4), ("ay", "ay_G", 4), ("az", "az_G", 4),
            ("gx", "gx_dps", 2), ("gy", "gy_dps", 2), ("gz", "gz_dps", 2),
            ("roll", "roll_deg", 2), ("pitch", "pitch_deg", 2))


def stamp(wall):
    return datetime.datetime.fromtimestamp(wall).strftime(TIME_FMT)[:-3]    # ms 까지


def _round(v, digits):
    return None if v is None else round(v, digits)


# --------------------------------------------------------------------------
# 기록
# --------------------------------------------------------------------------
class XlsxSink:
    """Frame + Result 를 CSV 한 행으로 쓰고, close() 에서 엑셀(측정 + 요약 시트)로 만든다."""

    def __init__(self, cfg, csv_path, xlsx_path):
        self.cfg = cfg
        self.csv_path, self.xlsx_path = csv_path, xlsx_path
        try:
            import openpyxl  # noqa: F401  끝나고 엑셀을 만들 때 쓴다. 없으면 미리 알려준다
            self.has_xlsx = True
        except ImportError:
            self.has_xlsx = False
            say("! openpyxl 이 없다 -> 끝나고 엑셀은 못 만들고 CSV 만 남는다.\n  %s\n" % PIP_HINT)

        self.weight, self.tension, self.corners = cfg.roles()
        self.checks = check_list(cfg)
        unit = "mm" if cfg.bed_l else "pct"
        cols = ["time", "t_s"]
        cols += ["%s_g" % l for l in self.weight]
        if self.weight:
            cols += ["total_g", "measured_kg"]
        if self.corners:
            cols += ["cog_x_" + unit, "cog_y_" + unit]
        cols += ["%s_N" % l for l in self.tension]
        # 판정 (analysis/analyzer.py)
        if self.weight:
            cols.append("load_kg")
        if self.corners:
            h = ("h_%s_mm" if cfg.bed_l else "k_%s_pct")
            cols += ["cog_pred_x_" + unit, "cog_pred_y_" + unit, h % "y", h % "x"]
        cols += ["%s_min_N" % l for l in self.tension]
        cols.append("risk")
        cols += ["lv_" + k for k, _ in self.checks]
        cols.append("risk_text")
        for l in cfg.imu_labels:
            cols += ["%s_%s" % (l, name) for _, name, _ in IMU_COLS]
        cols += ["%s_raw" % l for l in cfg.hx_labels]
        cols.append("flags")
        self.cols = cols

        self.f = open(csv_path, "w", newline="", encoding="utf-8-sig")
        self.w = csv.writer(self.f, lineterminator="\n")
        self.w.writerow(cols)
        self.rows = self.flagged = 0
        self.first_t = self.last_t = None
        self.last_flush = 0.0
        self.level_rows = {None: 0, OK: 0, CAUTION: 0, WARNING: 0, DANGER: 0}     # 종합 위험도별 행 수
        self.check_rows = {k: [0, 0, 0] for k, _ in self.checks}      # 판정별 [주의, 경고, 위험] 행 수
        self.events = []                        # 알림 레벨이 바뀐 순간 (엑셀 '판정' 시트)

    def row(self, frame, result):
        """Frame + Result -> CSV 한 행 (self.cols 순서)."""
        cfg = self.cfg
        c = result.cog
        out = [stamp(frame.wall), round(frame.t_s, 3)]
        out += [_round(frame.load_g[l], 1) for l in self.weight]
        if self.weight:
            out += [_round(result.total_g, 1), _round(result.measured_kg, 3)]
        if self.corners:
            out += [round(c.x, 1), round(c.y, 1)] if c else [None, None]
        out += [_round(frame.tension_n[l], 2) for l in self.tension]
        if self.weight:
            out.append(_round(result.load_kg, 3))
        if self.corners:
            tr = result.transfer
            if tr:
                out += [round(tr.pred_x, 1), round(tr.pred_y, 1)]
                out += [None if k is None else round(abs(k), 1) for k in (tr.ky, tr.kx)]
            else:
                out += [None] * 4
        out += [_round(result.tension_min.get(l), 2) for l in self.tension]
        out.append(result.level)
        out += [ch.level for ch in result.checks]
        out.append(" | ".join("%s %s: %s" % (risk_name(ch.level), ch.name, ch.text)
                              for ch in result.checks if ch.level))
        for l in cfg.imu_labels:
            r = frame.imu[l]
            if r is None:
                out += [None] * len(IMU_COLS)
            else:
                out += [round(getattr(r, attr), d) for attr, _, d in IMU_COLS]
        out += [frame.raw[l] for l in cfg.hx_labels]
        out.append("|".join(frame.flags))
        return out

    def write(self, frame, result):
        row = self.row(frame, result)
        self.w.writerow(row)
        self.rows += 1
        if row[-1]:
            self.flagged += 1
        self.level_rows[result.level] += 1
        for ch in result.checks:
            if ch.level:
                self.check_rows[ch.key][ch.level - 1] += 1
        for ch, old in result.changed:
            if len(self.events) < MAX_EVENTS:
                self.events.append((stamp(frame.wall), round(frame.t_s, 1), ch.name,
                                    risk_name(old), risk_name(ch.level), ch.text))
        if self.first_t is None:
            self.first_t = frame.t_s
        self.last_t = frame.t_s
        now = time.monotonic()
        if now - self.last_flush >= FLUSH_EVERY:
            self.last_flush = now
            self.f.flush()
            os.fsync(self.f.fileno())           # 전원이 나가도 1초 치만 잃게

    def close(self, src, notes):
        """CSV 를 닫고 엑셀을 만든다. src 는 FrameSource (요약 시트용). 돌려주는 값은 행 수."""
        self.f.close()
        if not self.rows:
            say("행이 하나도 없다. 배선/전원을 확인할 것.")
            return 0

        cfg = self.cfg
        span = self.last_t - self.first_t
        imu_err = ", ".join("%s %d회" % (l, n) for l, n in src.imus.errors.items() if n)
        info = [("모드", "시뮬레이션 (가짜 센서)" if src.is_sim else "실측"),
                ("HX711 타임아웃", src.timeouts),
                ("flags 있는 행", self.flagged)]
        if imu_err:
            info.append(("IMU 읽기 실패", imu_err))
        info += [("", ""), ("보정 파일", src.calib_src)]
        info += channel_info(cfg, src.hx.labels, src.calib, src.is_sim)
        if self.corners:
            c = self.corners
            info.append(("무게중심 원점", "%s 셀. X: %s->%s 방향, Y: %s->%s 방향, %s" % (
                c[0], c[0], c[1], c[0], c[2],
                "%g x %g mm" % (cfg.bed_l, cfg.bed_w) if cfg.bed_l
                else "% (50 = 가운데, BED_L/BED_W 를 주면 mm)")))
        info.append(("IMU 보정", src.imu_cal_src))
        where = [(l, hex(a)) for l, a in zip(cfg.imu_i2c_labels, cfg.imu_addrs)]
        if cfg.imu_udp_port:
            where.append((cfg.imu_udp_label, "ESP32 UDP %d" % cfg.imu_udp_port))
        for l, loc in where:
            if l not in src.imu_cal:
                desc = "못 엶 (빈칸)"
            else:
                ic = src.imu_cal[l]
                desc = ("up %s" % ic.get("up", "+z")) if ic else "보정 없음"
            info.append(("  " + l, "%s · %s" % (loc, desc)))
        info += self.judge_info()

        if self.has_xlsx:
            say("엑셀 만드는 중... (%d행)" % self.rows)
            build_xlsx(self.csv_path, self.xlsx_path, info, notes, extra=self.events_sheet)
            say("엑셀: %s" % self.xlsx_path)
        else:
            say("! openpyxl 이 없어 엑셀은 못 만들었다. CSV 는 엑셀로 바로 열린다.\n  %s" % PIP_HINT)
        say("CSV : %s" % self.csv_path)
        say("\n행 %d개  %.1f초  %.1fHz  타임아웃 %d  flags %d"
            % (self.rows, span, (self.rows - 1) / span if span > 0 else 0, src.timeouts, self.flagged))
        if src.timeouts or self.flagged:
            say("! 의심 행이 있다 (flags 열). sck_high 가 많으면 sudo -E 로 돌려볼 것.")
        lv = self.level_rows
        say("종합 위험도  정상 %d  주의 %d  경고 %d  위험 %d  판정 불가 %d (행)  알림 변화 %d번"
            % (lv[OK], lv[CAUTION], lv[WARNING], lv[DANGER], lv[None], len(self.events)))
        return self.rows

    def judge_info(self):
        """요약 시트용: 판정 기준과 위험도별 행 수."""
        lv = self.level_rows
        info = [("", "")]
        for i, line in enumerate(self.cfg.describe_judge().splitlines()):
            info.append(("판정 기준" if i == 0 else "", line.replace("판정", "", 1).strip()))
        info.append(("종합 위험도별 행 수", "정상 %d · 주의 %d · 경고 %d · 위험 %d · 판정 불가 %d"
                     % (lv[OK], lv[CAUTION], lv[WARNING], lv[DANGER], lv[None])))
        for k, nm in self.checks:
            caution, warning, danger = self.check_rows[k]
            info.append(("  " + nm, "주의 %d행 · 경고 %d행 · 위험 %d행" % (caution, warning, danger)))
        info.append(("알림 변화", "%d번 ('판정' 시트)%s" % (
            len(self.events), " · %d번에서 끊음" % MAX_EVENTS if len(self.events) >= MAX_EVENTS else "")))
        return info

    def events_sheet(self, wb, head_cell):
        """'판정' 시트: 판정별 알림 레벨이 바뀐 순간."""
        ws = wb.create_sheet("판정")
        for col, width in zip("ABCDEF", (24, 10, 12, 8, 8, 60)):
            ws.column_dimensions[col].width = width
        ws.append([head_cell(ws, t) for t in ("시각", "경과(초)", "판정", "이전", "지금", "내용")])
        for e in self.events:
            ws.append(list(e))


def channel_info(cfg, live, calib, is_sim):
    """요약 시트용: 설정된 채널마다 역할 / 핀 / 보정값. 안 보인 채널은 그렇다고 적는다."""
    _, tension, _ = cfg.roles()
    rows = []
    for i, l in enumerate(cfg.hx_labels):
        if is_sim:
            pins = "시뮬레이션"
        else:
            pins = "SCK GPIO%d / DT GPIO%d" % (
                cfg.hx_sck[0] if cfg.hx_shared_sck else cfg.hx_sck[i], cfg.hx_dt[i])
            if l in cfg.hx_invert:
                pins += " · 부호반전"
        off = calib["offset"].get(l)
        sc = calib["scale_g_per_count"].get(l)
        cal = ("offset %d, %.6g g/count" % (off, sc)
               if off is not None and sc is not None else "보정 없음")
        gone = "" if l in live else " · 시작할 때 응답 없음 (빈칸)"
        rows.append(("  " + l, "%s · %s · %s%s" % ("장력" if l in tension else "하중", pins, cal, gone)))
    return rows


# --------------------------------------------------------------------------
# 엑셀
# --------------------------------------------------------------------------
UNIT = {"g": "g", "N": "N", "G": "G", "dps": "°/s", "deg": "°", "mm": "mm", "pct": "%"}
FILL = {"time": "E7E6E6", "load": "DDEBF7", "tension": "FCE4D6",
        "imu": "E2EFDA", "judge": "FFF2CC", "diag": "EDEDED"}
CHECK_NAMES = {"load": "적재", "cog": "무게중심", "move": "화물 이동"}     # 그 밖의 lv_<라벨> 은 장력


def pretty(col):
    """CSV 열 이름 -> 엑셀 머리글."""
    fixed = {"time": "시각", "t_s": "경과(초)", "total_g": "총하중(g)", "flags": "flags",
             "measured_kg": "측정 무게(kg)", "load_kg": "적재 무게(kg, 장력 뺌)",
             "risk": "종합 위험도(0정상 1주의 2경고 3위험)", "risk_text": "위험 내용",
             "h_y_mm": "h 앞뒤(mm)", "h_x_mm": "h 좌우(mm)",
             "k_y_pct": "하중이동 앞뒤(%/G)", "k_x_pct": "하중이동 좌우(%/G)"}
    if col in fixed:
        return fixed[col]
    if col.startswith("lv_"):
        return "판정 " + CHECK_NAMES.get(col[3:], "장력 " + col[3:])
    name, _, unit = col.rpartition("_")
    if unit == "raw":
        return name + " raw"
    if name.startswith("cog_pred_"):
        return "예측 무게중심 %s(%s)" % (name[9:].upper(), UNIT[unit])
    if name.startswith("cog_"):
        return "무게중심 %s(%s)" % (name[4:].upper(), UNIT[unit])
    if name.endswith("_min") and unit == "N":
        return "%s 장력 하한(N)" % name[:-4]
    if unit == "N":
        return "%s 장력(N)" % name
    if unit in UNIT:
        return "%s(%s)" % (name.replace("_", " "), UNIT[unit])
    return col


def group(col):
    """머리글 색 묶음."""
    if col in ("time", "t_s"):
        return "time"
    if col in ("load_kg", "risk", "risk_text") or col.endswith("_min_N") \
            or col.startswith(("cog_pred_", "h_", "k_", "lv_")):
        return "judge"
    if col.endswith("_raw") or col == "flags":
        return "diag"
    if col.endswith("_N"):
        return "tension"
    if col.endswith(("_G", "_dps", "_deg")):
        return "imu"
    return "load"


class Stat:
    """평균/표준편차를 한 번 훑으며 계산 (Welford)."""

    def __init__(self):
        self.n = 0
        self.mean = 0.0
        self.m2 = 0.0
        self.min = self.max = self.first = self.last = None

    def add(self, x):
        self.n += 1
        d = x - self.mean
        self.mean += d / self.n
        self.m2 += d * (x - self.mean)
        if self.first is None:
            self.first = self.min = self.max = x
        self.min = min(self.min, x)
        self.max = max(self.max, x)
        self.last = x

    @property
    def sd(self):
        return (self.m2 / self.n) ** 0.5 if self.n > 1 else 0.0


def build_xlsx(csv_path, xlsx_path, info=(), notes=(), style=None, extra=None):
    """CSV -> 엑셀 (측정 + 요약 시트). 길게 기록해도 메모리를 안 먹게 write-only 로
    한 줄씩 흘려 쓴다. 기록 도중 죽어서 CSV 만 남았을 때도 이 함수로 다시 만들 수 있다.
    style 은 CSV 열 이름 -> (머리글, 머리글 색). 비우면 센서 엑셀 열 (카메라 엑셀은 log_camera.py).
    extra(wb, head_cell) 는 요약 시트 뒤에 시트를 더 붙인다 (카메라 엑셀의 트윈 기준 · 높이맵).
    돌려주는 값은 엑셀에 적은 행 수."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    notes = list(notes)
    style = style or (lambda h: (pretty(h), FILL[group(h)]))
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("측정")

    def head_cell(sheet, text, fill="DDE5F0"):
        c = WriteOnlyCell(sheet, value=text)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center")
        return c

    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f)
        head = next(rd)
        for i, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(i)].width = \
                {"time": 24, "risk_text": 60}.get(h) or max(9, text_width(style(h)[0]) + 2)
        ws.freeze_panes = "C2"                  # 머리글 줄과 시각 열 고정
        ws.append([head_cell(ws, *style(h)) for h in head])

        stats = [Stat() for _ in head]
        first = last = None
        n = broken = cut = 0
        for rec in rd:
            if len(rec) != len(head):           # 전원이 나가 잘린 마지막 줄 등
                broken += 1
                continue
            if n >= EXCEL_MAX_ROWS - 1:
                cut += 1
                continue
            row = []
            for i, v in enumerate(rec):
                if i == 0:
                    try:
                        tv = datetime.datetime.strptime(v, TIME_FMT)
                    except ValueError:
                        row.append(v)
                        continue
                    c = WriteOnlyCell(ws, value=tv)
                    c.number_format = "yyyy-mm-dd hh:mm:ss.000"
                    row.append(c)
                    first = first or tv
                    last = tv
                elif v == "":
                    row.append(None)
                else:
                    try:
                        x = int(v) if v.lstrip("-").isdigit() else float(v)
                    except ValueError:          # flags 같은 글자
                        row.append(v)
                        continue
                    stats[i].add(x)
                    row.append(x)
            ws.append(row)
            n += 1

    s = wb.create_sheet("요약")
    s.column_dimensions["A"].width = 26
    for col in "BCDEFGH":
        s.column_dimensions[col].width = 14

    span = (last - first).total_seconds() if first and last else 0.0
    zone = time.strftime("%Z")
    top = [
        ("측정 시작", "%s (%s)" % (first.strftime("%Y-%m-%d %H:%M:%S"), zone) if first else "-"),
        ("측정 끝", last.strftime("%Y-%m-%d %H:%M:%S") if last else "-"),
        ("길이(초)", round(span, 1)),
        ("행 수", n),
        ("평균 속도(Hz)", round((n - 1) / span, 2) if span > 0 else "-"),
    ]
    for k, v in top + list(info):
        s.append([k, v])

    if broken:
        notes.append("열 개수가 안 맞아 건너뛴 줄 %d개 (기록 중 전원이 나간 마지막 줄 등)" % broken)
    if cut:
        notes.append("엑셀 시트 한계(%d행)를 넘은 %d행은 CSV 에만 있다" % (EXCEL_MAX_ROWS, cut))
    if notes:
        s.append([])
        for i, t in enumerate(notes):
            s.append(["메모" if i == 0 else "", t])

    s.append([])
    s.append([head_cell(s, t) for t in
              ("항목", "평균", "최소", "최대", "표준편차", "처음", "마지막", "개수")])
    for h, st in zip(head, stats):
        if h in ("time", "t_s", "seq", "frame_ts", "risk") or h.startswith("lv_") or st.n == 0:
            continue
        s.append([style(h)[0], round(st.mean, 4), st.min, st.max, round(st.sd, 4),
                  st.first, st.last, st.n])

    if extra:
        extra(wb, head_cell)
    wb.save(xlsx_path)
    return n
