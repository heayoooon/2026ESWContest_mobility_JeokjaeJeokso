#!/usr/bin/env python3
"""바닥 로드셀(하중) + 벨트 장력 + IMU + 문 센서를 한 시계로 읽어 **시간순 엑셀** 로 남긴다.

    python3 log_xlsx.py                      # 파이. Ctrl+C 까지
    LOG_SECONDS=60 python3 log_xlsx.py       # 60초
    python log_xlsx.py --sim                 # PC(VS Code). 파이 없이 가짜 센서로

핀/라벨/역할은 전부 환경변수다 (config.py). 예) 바닥 4모서리 + 벨트 2개 + 문 센서
    HX_DT=5,6,13,19,26,20 HX_LABELS=FL,FR,RL,RR,T1,T2 DOOR_PIN=23 python3 log_xlsx.py

한 행 = HX711 샘플 한 번 (RATE 10SPS 면 약 10Hz). HX711 변환을 기다리는 동안 IMU 를
IMU_HZ 로 읽어 두었다가 그 구간 평균을 같은 행에 붙인다. 문 센서는 행마다 한 번 읽는다.
무게(g)/장력(N) 열은 calib.json 에 보정값이 있는 채널만 나온다 (calibrate.py 참고).
시작할 때 DT 가 안 떨어지는 모듈은 빼고 기록한다 (화면 경고 + 요약 시트에 남긴다).

만들어지는 파일 (LOG_DIR/)
    <이름>.csv    기록하면서 1초마다 디스크에 쓴다. 도중에 전원이 나가도 여기까지 남고
                  엑셀로 바로 열린다 (UTF-8 BOM 이라 한글이 안 깨진다)
    <이름>.xlsx   끝날 때(Ctrl+C / LOG_SECONDS / kill) CSV 로 만든다. 측정 + 요약 시트
"""
import csv
import datetime
import os
import shutil
import signal
import sys
import time
import types
import unicodedata

from config import Config
from hx711_bus import HX711Timeout
from mpu6050 import cal_key, find_cal, load_cal, tilt, to_units

G_TO_N = 9.80665 / 1000     # 질량(g) -> 힘(N)
COG_MIN_G = 100             # 총하중이 이보다 작으면 무게중심은 노이즈뿐이라 비운다
ALPHA = 0.98                # 기울기 상보필터 (pi/imu/imu_angle.py 와 같은 값)
STATUS_EVERY = 0.5          # 화면 갱신(초)
FLUSH_EVERY = 1.0           # CSV 를 디스크까지 내리는 주기(초)
EXCEL_MAX_ROWS = 1048576    # 엑셀 시트 하나의 행 한계
TIME_FMT = "%Y-%m-%d %H:%M:%S.%f"
PIP_HINT = "pip install openpyxl   (파이: sudo apt install python3-openpyxl)"


# --------------------------------------------------------------------------
# 출력
# --------------------------------------------------------------------------
def say(*args):
    """SSH 가 끊겨 터미널이 사라져도 기록은 계속되게, 출력 오류는 삼킨다."""
    try:
        print(*args, flush=True)
    except OSError:
        pass


def text_width(s):
    """터미널/엑셀에서 차지하는 칸 수. 한글은 2칸."""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in s)


def status(text):
    """한 줄 상태표시. 터미널 폭을 넘기면 줄이 밀려 \\r 이 안 먹으니 잘라서 쓴다."""
    cols = shutil.get_terminal_size((120, 20)).columns - 1
    out, w = "", 0
    for ch in text:
        cw = text_width(ch)
        if w + cw > cols:
            break
        out += ch
        w += cw
    try:
        sys.stdout.write("\r" + out + " " * (cols - w))
        sys.stdout.flush()
    except OSError:
        pass


def boost_priority():
    """비트뱅잉이 스케줄러에 밀리면 SCK 가 오래 HIGH 로 떠서 칩이 파워다운된다."""
    try:
        os.sched_setscheduler(0, os.SCHED_FIFO, os.sched_param(10))
        return "SCHED_FIFO(10)"
    except Exception:
        pass
    try:
        os.nice(-10)
        return "nice -10"
    except Exception:
        return "일반 우선순위 (sudo -E 로 돌리면 타이밍이 안정된다)"


def stamp(wall):
    return datetime.datetime.fromtimestamp(wall).strftime(TIME_FMT)[:-3]    # ms 까지


# --------------------------------------------------------------------------
# IMU: HX711 을 기다리는 동안 읽어서 행마다 평균
# --------------------------------------------------------------------------
class ImuSet:
    """IMU 들을 IMU_HZ 로 읽어 쌓아두고, 행마다 구간 평균 + 현재 기울기를 내준다."""
    FIELDS = ("ax_G", "ay_G", "az_G", "gx_dps", "gy_dps", "gz_dps", "roll_deg", "pitch_deg")
    DIGITS = (4, 4, 4, 2, 2, 2, 2, 2)

    def __init__(self, sensors, hz):
        self.sensors = sensors                  # [(라벨, 장치, imu_cal 항목 또는 None)]
        self.labels = [s[0] for s in sensors]
        self.period = 1.0 / max(1, hz)
        self.next_t = time.monotonic()
        self.sum = {l: [0.0] * 6 for l in self.labels}
        self.n = {l: 0 for l in self.labels}
        self.angle = {l: None for l in self.labels}     # (roll, pitch, 마지막 시각)
        self.errors = {l: 0 for l in self.labels}

    def poll(self):
        """주기가 됐으면 전 IMU 를 한 번씩 읽는다. 아니면 바로 돌아간다."""
        now = time.monotonic()
        if not self.sensors or now < self.next_t:
            return
        self.next_t += self.period
        if self.next_t < now:
            self.next_t = now + self.period     # 밀렸으면 기준을 다시 잡는다
        for label, dev, cal in self.sensors:
            try:
                r = dev.read_raw()
            except Exception:                   # I2C 한 번 실패는 흔하다. 세고 넘어간다
                self.errors[label] += 1
                continue
            v = to_units(r[1:7], cal)
            s = self.sum[label]
            for i in range(6):
                s[i] += v[i]
            self.n[label] += 1

            # 상보필터: 자이로를 적분하고 가속도(중력 방향)로 천천히 끌어당긴다
            ar, ap = tilt(v[0], v[1], v[2])
            t = r[0] / 1e9
            prev = self.angle[label]
            if prev is None:
                self.angle[label] = (ar, ap, t)
            else:
                dt = t - prev[2]
                self.angle[label] = (ALPHA * (prev[0] + v[3] * dt) + (1 - ALPHA) * ar,
                                     ALPHA * (prev[1] + v[4] * dt) + (1 - ALPHA) * ap, t)

    def take(self):
        """지난 take() 이후 평균. 그 사이 한 번도 못 읽은 IMU 는 None."""
        out = {}
        for l in self.labels:
            n = self.n[l]
            if n == 0:
                out[l] = None
                continue
            vals = [x / n for x in self.sum[l]] + list(self.angle[l][:2])
            out[l] = [round(v, d) for v, d in zip(vals, self.DIGITS)]
            self.sum[l] = [0.0] * 6
            self.n[l] = 0
        return out

    def close(self):
        for _, dev, _ in self.sensors:
            dev.close()


# --------------------------------------------------------------------------
# 열 구성과 한 행 계산
# --------------------------------------------------------------------------
class Table:
    """채널 역할(하중/장력)과 보정 여부로 열을 정하고, 샘플 하나를 한 행으로 바꾼다."""

    def __init__(self, cfg, labels, cal, imu_labels, has_door):
        weight, tension, corners = cfg.roles(labels)
        self.labels = labels
        self.cal = cal                          # {라벨: (offset, g/count)}
        self.wg = [l for l in weight if l in cal]
        self.tn = [l for l in tension if l in cal]
        self.total = bool(weight) and len(self.wg) == len(weight)   # 전부 보정됐을 때만 합계
        self.corners = corners if corners and all(c in cal for c in corners) else None
        if cfg.bed_l:
            self.bed, unit, self.cog_unit = (cfg.bed_l, cfg.bed_w), "mm", "mm"
        else:
            self.bed, unit, self.cog_unit = (100.0, 100.0), "pct", "%"
        self.imu_labels = imu_labels
        self.door = has_door

        cols = ["time", "t_s"]
        cols += ["%s_g" % l for l in self.wg]
        if self.total:
            cols.append("total_g")
        if self.corners:
            cols += ["cog_x_" + unit, "cog_y_" + unit]
        cols += ["%s_N" % l for l in self.tn]
        for l in imu_labels:
            cols += ["%s_%s" % (l, k) for k in ImuSet.FIELDS]
        if has_door:
            cols.append("door_open")
        cols += ["%s_raw" % l for l in labels]
        cols.append("flags")
        self.cols = cols

    def row(self, wall, t_rel, s, imu, door_open):
        """(CSV 한 행, 화면에 띄울 요약)"""
        raw = dict(zip(self.labels, s.values))
        g = {l: (raw[l] - off) * sc for l, (off, sc) in self.cal.items() if l in raw}
        flags = list(s.flags)
        out = [stamp(wall), round(t_rel, 3)]
        out += [round(g[l], 1) for l in self.wg]

        total = None
        if self.total:
            total = sum(g[l] for l in self.wg)
            out.append(round(total, 1))

        cog = None
        if self.corners:
            # 계획서 식: X = (FR+RR)/M * L,  Y = (RL+RR)/M * W  (앞왼 셀이 원점)
            fl, fr, rl, rr = (g[c] for c in self.corners)
            m = fl + fr + rl + rr
            if m >= COG_MIN_G:
                cog = ((fr + rr) / m * self.bed[0], (rl + rr) / m * self.bed[1])
                out += [round(cog[0], 1), round(cog[1], 1)]
            else:
                out += [None, None]

        tension = [g[l] * G_TO_N for l in self.tn]
        out += [round(v, 2) for v in tension]

        for l in self.imu_labels:
            v = imu.get(l)
            if v is None:
                out += [None] * len(ImuSet.FIELDS)
                flags.append("imu_miss:" + l)
            else:
                out += v
        if self.door:
            out.append(door_open)
        out += s.values
        out.append("|".join(flags))

        show = []
        if total is not None:
            show.append("하중 %.0fg" % total)
        if cog:
            show.append("CoG %.0f,%.0f%s" % (cog[0], cog[1], self.cog_unit))
        show += ["%s %.1fN" % (l, v) for l, v in zip(self.tn, tension)]
        if total is None and not self.tn:       # 보정이 없으면 raw 라도
            show += ["%s=%d" % (l, v) for l, v in zip(self.labels, s.values)]
        if self.imu_labels and imu.get(self.imu_labels[0]):
            v = imu[self.imu_labels[0]]
            show.append("%s ax%+.2f ay%+.2f" % (self.imu_labels[0], v[0], v[1]))
        if self.door:
            show.append("문 열림" if door_open else "문 닫힘")
        return out, "  ".join(show)


# --------------------------------------------------------------------------
# 장치 열기
# --------------------------------------------------------------------------
def open_hx(cfg, keep=None):
    """HX711Bus 를 연다. keep 을 주면 그 인덱스의 모듈만."""
    from hx711_bus import HX711Bus
    idx = list(range(len(cfg.hx_dt))) if keep is None else keep
    sck = cfg.hx_sck[0] if cfg.hx_shared_sck else [cfg.hx_sck[i] for i in idx]
    return HX711Bus(sck, [cfg.hx_dt[i] for i in idx], [cfg.hx_labels[i] for i in idx],
                    gain=cfg.hx_gain, pull=cfg.hx_pull, chip=cfg.hx_chip,
                    invert=cfg.hx_invert)


def open_real(cfg, notes):
    from calibrate import CALIB_PATH, load_calib
    from mpu6050 import MPU6050

    try:
        hx = open_hx(cfg)
    except Exception as e:
        sys.exit("HX711 열기 실패: %s: %s\nPC 에서 시험하려면: python log_xlsx.py --sim"
                 % (type(e).__name__, e))
    hx.sync_reset()                             # 전 모듈 변환 위상 정렬
    missing = hx.missing()
    if missing:
        if len(missing) == hx.n:
            hx.close()
            sys.exit("HX711 이 하나도 안 보인다 (DT 가 계속 HIGH). "
                     "전원/배선부터 확인: python3 check_wiring.py")
        say("! 안 보이는 모듈 %s -> 빼고 기록한다 (DT 가 계속 HIGH)" % ", ".join(missing))
        notes.append("시작할 때 안 보여서 뺀 채널: %s" % ", ".join(missing))
        keep = [i for i, l in enumerate(cfg.hx_labels) if l not in missing]
        hx.close()
        hx = open_hx(cfg, keep)
        hx.sync_reset()

    cal_path = find_cal(cfg.imu_cal)
    imu_cal = load_cal(cal_path) if cal_path else {}
    imus = []
    for addr, label in zip(cfg.imu_addrs, cfg.imu_labels):
        m = None
        try:
            m = MPU6050(addr, bus=cfg.imu_bus, hz=cfg.imu_hz, label=label)
            m.check()
        except Exception as e:
            if m is not None:
                m.close()
            say("! IMU %s(%s) 못 엶: %s" % (label, hex(addr), e))
            notes.append("IMU %s(%s) 못 엶: %s" % (label, hex(addr), e))
            continue
        imus.append((label, m, imu_cal.get(cal_key(cfg.imu_bus, addr))))
    if not imus:
        say("  i2cdetect -y %d 로 68/69 가 보이는지. 진단 스크립트를 돌린 뒤라면 "
            "pinctrl set 2 a3 pu; pinctrl set 3 a3 pu 로 I2C 핀을 되돌린다" % cfg.imu_bus)

    door = None
    if cfg.door_pin is not None:
        from door import Door
        try:
            door = Door(cfg.door_pin, cfg.door_invert, cfg.hx_chip)
        except Exception as e:
            say("! 문 센서 GPIO%d 못 엶: %s" % (cfg.door_pin, e))
            notes.append("문 센서 GPIO%d 못 엶: %s" % (cfg.door_pin, e))

    calib = load_calib()
    return types.SimpleNamespace(
        hx=hx, imus=imus, door=door, calib=calib,
        calib_src=("%s (%s)" % (CALIB_PATH, calib.get("updated", "?"))
                   if os.path.exists(CALIB_PATH) else "없음"),
        imu_cal_src=cal_path or ("IMU_CAL=%s 파일 없음" % cfg.imu_cal if cfg.imu_cal
                                 else "없음 (바이어스/축 보정 안 함)"))


def open_sim(cfg):
    import sim
    scn = sim.Scenario()
    hx = sim.SimHX711Bus(cfg, scn)
    imus = [(l, sim.SimMPU6050(l, scn, i), None) for i, l in enumerate(cfg.imu_labels)]
    return types.SimpleNamespace(
        hx=hx, imus=imus, door=sim.SimDoor(scn), calib=hx.calib(),
        calib_src="시뮬레이션 내장값", imu_cal_src="시뮬레이션 (평평하게 놓임)")


def channel_info(cfg, labels, calib, is_sim):
    """요약 시트용: 채널마다 역할 / 핀 / 보정값."""
    _, tension, _ = cfg.roles(labels)
    rows = []
    for l in labels:
        i = cfg.hx_labels.index(l)
        if is_sim:
            pins = "시뮬레이션"
        else:
            pins = "SCK GPIO%d / DT GPIO%d" % (
                cfg.hx_sck[0] if cfg.hx_shared_sck else cfg.hx_sck[i], cfg.hx_dt[i])
            if l in cfg.hx_invert:
                pins += " · 부호반전"
        off = calib["offset"].get(l)
        sc = calib["scale_g_per_count"].get(l)
        cal = ("offset %d, %.6g g/count" % (off, sc)
               if off is not None and sc is not None else "보정 없음")
        rows.append(("  " + l, "%s · %s · %s" % ("장력" if l in tension else "하중", pins, cal)))
    return rows


# --------------------------------------------------------------------------
# 엑셀
# --------------------------------------------------------------------------
UNIT = {"g": "g", "N": "N", "G": "G", "dps": "°/s", "deg": "°", "mm": "mm", "pct": "%"}
FILL = {"time": "E7E6E6", "load": "DDEBF7", "tension": "FCE4D6",
        "imu": "E2EFDA", "door": "FFF2CC", "diag": "EDEDED"}


def pretty(col):
    """CSV 열 이름 -> 엑셀 머리글."""
    fixed = {"time": "시각", "t_s": "경과(초)", "total_g": "총하중(g)",
             "door_open": "문 열림(1=열림)", "flags": "flags"}
    if col in fixed:
        return fixed[col]
    name, _, unit = col.rpartition("_")
    if unit == "raw":
        return name + " raw"
    if name.startswith("cog_"):
        return "무게중심 %s(%s)" % (name[4:].upper(), UNIT[unit])
    if unit == "N":
        return "%s 장력(N)" % name
    if unit in UNIT:
        return "%s(%s)" % (name.replace("_", " "), UNIT[unit])
    return col


def group(col):
    """머리글 색 묶음."""
    if col in ("time", "t_s"):
        return "time"
    if col.endswith("_raw") or col == "flags":
        return "diag"
    if col.endswith("_N"):
        return "tension"
    if col == "door_open":
        return "door"
    if col.endswith(("_G", "_dps", "_deg")):
        return "imu"
    return "load"


class Stat:
    """평균/표준편차를 한 번 훑으며 계산 (Welford)."""

    def __init__(self):
        self.n = 0
        self.mean = 0.0
        self.m2 = 0.0
        self.min = self.max = self.first = self.last = None

    def add(self, x):
        self.n += 1
        d = x - self.mean
        self.mean += d / self.n
        self.m2 += d * (x - self.mean)
        if self.first is None:
            self.first = self.min = self.max = x
        self.min = min(self.min, x)
        self.max = max(self.max, x)
        self.last = x

    @property
    def sd(self):
        return (self.m2 / self.n) ** 0.5 if self.n > 1 else 0.0


def build_xlsx(csv_path, xlsx_path, info=(), notes=()):
    """CSV -> 엑셀 (측정 + 요약 시트). 길게 기록해도 메모리를 안 먹게 write-only 로
    한 줄씩 흘려 쓴다. 기록 도중 죽어서 CSV 만 남았을 때도 이 함수로 다시 만들 수 있다.
    돌려주는 값은 엑셀에 적은 행 수."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    notes = list(notes)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("측정")

    def head_cell(sheet, text, fill="DDE5F0"):
        c = WriteOnlyCell(sheet, value=text)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center")
        return c

    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f)
        head = next(rd)
        for i, h in enumerate(head, 1):
            ws.column_dimensions[get_column_letter(i)].width = \
                24 if h == "time" else max(9, text_width(pretty(h)) + 2)
        ws.freeze_panes = "C2"                  # 머리글 줄과 시각 열 고정
        ws.append([head_cell(ws, pretty(h), FILL[group(h)]) for h in head])

        stats = [Stat() for _ in head]
        first = last = None
        n = broken = cut = 0
        for rec in rd:
            if len(rec) != len(head):           # 전원이 나가 잘린 마지막 줄 등
                broken += 1
                continue
            if n >= EXCEL_MAX_ROWS - 1:
                cut += 1
                continue
            row = []
            for i, v in enumerate(rec):
                if i == 0:
                    try:
                        tv = datetime.datetime.strptime(v, TIME_FMT)
                    except ValueError:
                        row.append(v)
                        continue
                    c = WriteOnlyCell(ws, value=tv)
                    c.number_format = "yyyy-mm-dd hh:mm:ss.000"
                    row.append(c)
                    first = first or tv
                    last = tv
                elif v == "":
                    row.append(None)
                else:
                    try:
                        x = int(v) if v.lstrip("-").isdigit() else float(v)
                    except ValueError:          # flags 같은 글자
                        row.append(v)
                        continue
                    stats[i].add(x)
                    row.append(x)
            ws.append(row)
            n += 1

    s = wb.create_sheet("요약")
    s.column_dimensions["A"].width = 26
    for col in "BCDEFGH":
        s.column_dimensions[col].width = 14

    span = (last - first).total_seconds() if first and last else 0.0
    zone = time.strftime("%Z")
    top = [
        ("측정 시작", "%s (%s)" % (first.strftime("%Y-%m-%d %H:%M:%S"), zone) if first else "-"),
        ("측정 끝", last.strftime("%Y-%m-%d %H:%M:%S") if last else "-"),
        ("길이(초)", round(span, 1)),
        ("행 수", n),
        ("평균 속도(Hz)", round((n - 1) / span, 2) if span > 0 else "-"),
    ]
    for k, v in top + list(info):
        s.append([k, v])

    if broken:
        notes.append("열 개수가 안 맞아 건너뛴 줄 %d개 (기록 중 전원이 나간 마지막 줄 등)" % broken)
    if cut:
        notes.append("엑셀 시트 한계(%d행)를 넘은 %d행은 CSV 에만 있다" % (EXCEL_MAX_ROWS, cut))
    if notes:
        s.append([])
        for i, t in enumerate(notes):
            s.append(["메모" if i == 0 else "", t])

    s.append([])
    s.append([head_cell(s, t) for t in
              ("항목", "평균", "최소", "최대", "표준편차", "처음", "마지막", "개수")])
    for h, st in zip(head, stats):
        if h in ("time", "t_s") or st.n == 0:
            continue
        s.append([pretty(h), round(st.mean, 4), st.min, st.max, round(st.sd, 4),
                  st.first, st.last, st.n])

    wb.save(xlsx_path)
    return n


# --------------------------------------------------------------------------
# 메인
# --------------------------------------------------------------------------
def main():
    args = sys.argv[1:]
    if any(a != "--sim" for a in args):
        sys.exit(__doc__)
    is_sim = "--sim" in args
    if is_sim and not os.environ.get("HX_LABELS") and not os.environ.get("HX_DT"):
        os.environ["HX_LABELS"] = "FL,FR,RL,RR,T1,T2"   # 가짜 센서는 최종 배치(4모서리+벨트2)로

    cfg = Config()
    name = cfg.log_name
    notes = []
    if is_sim:
        if not os.environ.get("LOG_NAME"):
            name = "sim_" + name
        say("*** 시뮬레이션: 가짜 센서 값이다. 파이에서 실측할 때는 --sim 을 빼라 ***")
        notes.append("시뮬레이션 데이터 (가짜 센서). 실측 아님.")
    say(cfg.describe())
    say()
    try:
        import openpyxl  # noqa: F401  끝나고 엑셀을 만들 때 쓴다. 없으면 미리 알려준다
        has_xlsx = True
    except ImportError:
        has_xlsx = False
        say("! openpyxl 이 없다 -> 끝나고 엑셀은 못 만들고 CSV 만 남는다.\n  %s\n" % PIP_HINT)

    dev = open_sim(cfg) if is_sim else open_real(cfg, notes)
    hx = dev.hx
    say(hx.describe())
    flipped = []
    if not is_sim:
        from calibrate import stale
        flipped = stale(dev.calib, cfg, hx.labels)
    if flipped:
        say("! HX_INVERT 가 영점 잡을 때와 다른 채널 %s -> raw 만 적는다. calibrate.py tare 부터 다시"
            % ", ".join(flipped))
        notes.append("HX_INVERT 가 영점 잡을 때와 달라 보정을 안 쓴 채널: %s" % ", ".join(flipped))
    cal = {l: (dev.calib["offset"][l], dev.calib["scale_g_per_count"][l])
           for l in hx.labels
           if l not in flipped
           and dev.calib["offset"].get(l) is not None
           and dev.calib["scale_g_per_count"].get(l) is not None}
    if not cal:
        say("보정  없음 -> 원시 카운트만 적는다. calibrate.py 로 tare / span / cell 을 먼저 하면 "
            "무게(g)·장력(N) 열이 생긴다.")
        notes.append("calib.json 에 보정값이 없어 무게(g)/장력(N) 없이 raw 만 적었다.")
    elif len(cal) < len(hx.labels):
        miss = [l for l in hx.labels if l not in cal]
        say("보정  %d/%d 채널 (없는 채널: %s -> raw 만)" % (len(cal), len(hx.labels), ", ".join(miss)))
        notes.append("보정값이 없어 raw 만 적은 채널: %s" % ", ".join(miss))

    imus = ImuSet(dev.imus, cfg.imu_hz)
    door = dev.door
    table = Table(cfg, hx.labels, cal, imus.labels, door is not None)
    for l, _, c in dev.imus:
        say("IMU   %s 열림  보정 %s" % (l, ("있음 (up %s)" % c.get("up", "+z")) if c else "없음"))
    if not is_sim:
        say("우선순위  %s" % boost_priority())

    os.makedirs(cfg.log_dir, exist_ok=True)
    csv_path = os.path.join(cfg.log_dir, name + ".csv")
    xlsx_path = os.path.join(cfg.log_dir, name + ".xlsx")

    stop = {"now": False}

    def on_stop(*_):
        stop["now"] = True

    # Ctrl+C / kill / SSH 끊김 모두 '멈추고 엑셀 만들기'. nohup 처럼 무시로 걸려 있으면 그대로 둔다
    for sig_name in ("SIGINT", "SIGTERM", "SIGHUP"):
        sig = getattr(signal, sig_name, None)
        if sig is not None and signal.getsignal(sig) is not signal.SIG_IGN:
            signal.signal(sig, on_stop)

    try:
        hx.read(timeout=cfg.hx_timeout)         # 파워다운 직후 첫 값은 버린다
    except Exception:
        pass

    t0_mono = time.monotonic_ns()
    t0_wall = time.time()
    deadline = None if cfg.log_seconds <= 0 else t0_mono + int(cfg.log_seconds * 1e9)
    rows = timeouts = flagged = 0
    first_t = last_t = None
    last_status = last_flush = 0.0
    say("\n기록 시작 -> %s\nCtrl+C 로 끝내면 엑셀을 만든다.\n" % csv_path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(table.cols)
        try:
            while not stop["now"]:
                if deadline and time.monotonic_ns() >= deadline:
                    break
                # HX711 변환(~100ms)을 기다리는 동안 IMU 를 IMU_HZ 로 읽는다
                limit = time.monotonic() + cfg.hx_timeout
                while not hx.ready() and time.monotonic() < limit and not stop["now"]:
                    imus.poll()
                    time.sleep(0.0005)
                if stop["now"]:
                    break
                try:
                    s = hx.read(timeout=0.05)   # 준비됐으면 바로 클럭, 아니면 곧 타임아웃
                except HX711Timeout as e:
                    timeouts += 1
                    say("\n타임아웃 (DRDY 안 옴): %s" % ", ".join(e.waiting))
                    hx.sync_reset(settle=0.3)   # 위상 다시 맞추고 계속
                    continue

                t_rel = (s.t_ns - t0_mono) / 1e9
                row, shown = table.row(t0_wall + t_rel, t_rel, s, imus.take(),
                                       door.is_open() if door else None)
                w.writerow(row)
                rows += 1
                if row[-1]:
                    flagged += 1
                if first_t is None:
                    first_t = t_rel
                last_t = t_rel

                now = time.monotonic()
                if now - last_flush >= FLUSH_EVERY:
                    last_flush = now
                    f.flush()
                    os.fsync(f.fileno())        # 전원이 나가도 1초 치만 잃게
                if now - last_status >= STATUS_EVERY:
                    last_status = now
                    status("%6.1fs %6d행  %s%s" % (t_rel, rows, shown,
                                                   "  타임아웃 %d" % timeouts if timeouts else ""))
        except Exception as e:                  # 예상 못 한 오류라도 지금까지 것은 엑셀로 남긴다
            say("\n! 오류로 기록 중단: %s: %s" % (type(e).__name__, e))
            notes.append("오류로 기록 중단: %s: %s" % (type(e).__name__, e))
        finally:
            hx.close()
            imus.close()
            if door:
                door.close()

    say("\n\n정리 중...")
    if not rows:
        say("행이 하나도 없다. 배선/전원을 확인할 것.")
        return 1

    span = last_t - first_t
    imu_err = ", ".join("%s %d회" % (l, n) for l, n in imus.errors.items() if n)
    info = [("모드", "시뮬레이션 (가짜 센서)" if is_sim else "실측"),
            ("HX711 타임아웃", timeouts),
            ("flags 있는 행", flagged)]
    if imu_err:
        info.append(("IMU 읽기 실패", imu_err))
    info += [("", ""), ("보정 파일", dev.calib_src)]
    info += channel_info(cfg, hx.labels, dev.calib, is_sim)
    if table.corners:
        info.append(("무게중심 원점", "%s 셀. X: %s->%s 방향, Y: %s->%s 방향, %s" % (
            table.corners[0], table.corners[0], table.corners[1],
            table.corners[0], table.corners[2],
            "%g x %g mm" % table.bed if cfg.bed_l else "% (50 = 가운데, BED_L/BED_W 를 주면 mm)")))
    info.append(("IMU 보정", dev.imu_cal_src))
    for l, _, c in dev.imus:
        addr = cfg.imu_addrs[cfg.imu_labels.index(l)]
        info.append(("  " + l, "%s · %s" % (hex(addr), ("up %s" % c.get("up", "+z")) if c else "보정 없음")))
    info.append(("문 센서", door.describe() + ", 1 = 열림" if door else "없음"))

    if has_xlsx:
        say("엑셀 만드는 중... (%d행)" % rows)
        build_xlsx(csv_path, xlsx_path, info, notes)
        say("엑셀: %s" % xlsx_path)
    else:
        say("! openpyxl 이 없어 엑셀은 못 만들었다. CSV 는 엑셀로 바로 열린다.\n  %s" % PIP_HINT)
    say("CSV : %s" % csv_path)
    say("\n행 %d개  %.1f초  %.1fHz  타임아웃 %d  flags %d"
        % (rows, span, (rows - 1) / span if span > 0 else 0, timeouts, flagged))
    if timeouts or flagged:
        say("! 의심 행이 있다 (flags 열). sck_high 가 많으면 sudo -E 로 돌려볼 것.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
